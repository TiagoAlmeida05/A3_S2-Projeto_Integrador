import app.models as models

from docx import Document
from docx.shared import Pt
import io
import urllib.parse

def create_codebook_docx(db, project):
    codes = project.codes

    code_ids = [c.id for c in codes]
    memos = db.query(models.Memo).filter(
        models.Memo.target_type == "code",
        models.Memo.target_id.in_(code_ids)
    ).all()

    # group memos by code_id
    memo_dict = {}
    for m in memos:
        if m.target_id not in memo_dict:
            memo_dict[m.target_id] = []
        memo_dict[m.target_id].append(m.text)

    # build the hierarchical tree recursively
    def build_tree(parent_id=None, depth=0):
        tree = []
        for c in [c for c in codes if c.parent_id == parent_id]:
            tree.append((c, depth))
            tree.extend(build_tree(c.id, depth + 1))
        return tree

    ordered_codes = build_tree()

    # initialize the Word Document
    doc = Document()
    doc.add_heading(f"Codebook: {project.name}", 0)
    doc.add_paragraph(f"Exported from jUPiter QDA on {project.last_accessed.strftime('%B %d, %Y')}")

    # populate the document
    for code, depth in ordered_codes:
        # heading levels 1 to 4
        level = min(depth + 1, 4)
        heading = doc.add_heading(level=level)

        # create an indent string (4 spaces per depth level)
        indent_prefix = "    " * depth

        # add the spaces before the code name
        run = heading.add_run(f"{indent_prefix}{code.name}")

        if code.id in memo_dict:
            for memo_text in memo_dict[code.id]:
                p = doc.add_paragraph(style='List Bullet')
                # use Word's native left indent for bullets so the actual bullet dot moves over!
                p.paragraph_format.left_indent = Pt(24 * (depth + 1))
                p.add_run("Memo: ").bold = True
                p.add_run(memo_text)

    # save to a virtual file in memory
    mem_stream = io.BytesIO()
    doc.save(mem_stream)
    mem_stream.seek(0)

    # safely encode the filename
    safe_filename = urllib.parse.quote(f"Codebook_{project.name}.docx")
    return mem_stream,safe_filename