import io
import uuid
import xml.etree.ElementTree as ET
import zipfile

from fastapi import HTTPException, UploadFile, Depends, File
from fastapi.responses import Response
from sqlalchemy.orm import Session
from database import engine, get_db
from datetime import datetime

import models

def parse_refi_date(date_str):
    if not date_str:
        return None
    try:
        return datetime.fromisoformat(date_str.replace("Z", "+00:00"))
    except ValueError:
        return None
    
def generate_guid(prefix: str, item_id: int) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_DNS, f"jupiter.qda.{prefix}.{item_id}"))


def export_refi_xml(project_id: int, db: Session = Depends(get_db)):
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    ET.register_namespace("", "urn:QDA-XML:project:1.0")
    
    # Generate a single user GUID to own the project
    master_user_guid = str(uuid.uuid4())

    #Root Element
    root = ET.Element("{urn:QDA-XML:project:1.0}Project", attrib={
        "name": project.name,
        "origin": "jUPiter QDA",
        "creatingUserGUID": master_user_guid
    })

    # Users
    users = ET.SubElement(root, "{urn:QDA-XML:project:1.0}Users")
    ET.SubElement(users, "{urn:QDA-XML:project:1.0}User", attrib={
        "guid": master_user_guid, 
        "name": "jUPiter User"
    })

    codes = db.query(models.Code).filter(models.Code.project_id == project_id).all()
    docs = db.query(models.Document).filter(models.Document.project_id == project_id).all()
    doc_segments = db.query(models.Segment).join(models.Document).filter(models.Document.project_id == project_id).all()
    
    code_ids = [c.id for c in codes]
    segment_ids = [s.id for s in doc_segments]

    # Fetch Memos by target type
    project_memos = db.query(models.Memo).filter(models.Memo.target_type == "project", models.Memo.target_id == project_id).all()
    code_memos = db.query(models.Memo).filter(models.Memo.target_type == "code", models.Memo.target_id.in_(code_ids)).all() if code_ids else []
    segment_memos = db.query(models.Memo).filter(models.Memo.target_type == "segment", models.Memo.target_id.in_(segment_ids)).all() if segment_ids else []
    
    all_memos = project_memos + code_memos + segment_memos
    
    memos_by_code = {}
    for m in code_memos:
        memos_by_code.setdefault(m.target_id, []).append(m)

    memos_by_segment = {}
    for m in segment_memos:
        memos_by_segment.setdefault(m.target_id, []).append(m)
        
    codes_by_parent = {}
    for c in codes:
        codes_by_parent.setdefault(c.parent_id, []).append(c)

    # CodeBook 
    codebook = ET.SubElement(root, "{urn:QDA-XML:project:1.0}CodeBook")
    codes_elem = ET.SubElement(codebook, "{urn:QDA-XML:project:1.0}Codes")

    code_guid_map = {}

    def build_code_xml(c, parent_xml_element):
        cg = generate_guid("code", c.id)
        code_guid_map[c.id] = cg
        code_attribs = {"guid": cg, "name": c.name, "isCodable": "true"}
        if c.color: 
            code_attribs["color"] = c.color
            
        code_elem = ET.SubElement(parent_xml_element, "{urn:QDA-XML:project:1.0}Code", attrib=code_attribs)


        if c.id in memos_by_code:
            for m in memos_by_code[c.id]:
                ET.SubElement(code_elem, "{urn:QDA-XML:project:1.0}NoteRef", attrib={
                    "targetGUID": generate_guid("memo", m.id)
                })
            
        child_codes = codes_by_parent.get(c.id, [])
        for child in child_codes:
            build_code_xml(child, code_elem)

    #start with recursion from parent codes
    top_level_codes = codes_by_parent.get(None, [])
    for tlc in top_level_codes:
        build_code_xml(tlc, codes_elem)

    zip_files_to_write = {}

    # Sources
    sources = ET.SubElement(root, "{urn:QDA-XML:project:1.0}Sources")
   
    for d in docs:
        doc_guid = generate_guid("doc", d.id)
        internal_filename = f"{doc_guid}.txt"
        
        source_elem = ET.SubElement(sources, "{urn:QDA-XML:project:1.0}TextSource", attrib={
            "guid": doc_guid,
            "name": d.filename,
            "plainTextPath": f"internal://{internal_filename}",
            "creatingUser": master_user_guid
        })

        zip_files_to_write[f"Sources/{internal_filename}"] = d.content

        d_segments = [s for s in doc_segments if s.document_id == d.id]
        for seg in d_segments:
            sel_guid = generate_guid("selection", seg.id)
            
            sel_elem = ET.SubElement(source_elem, "{urn:QDA-XML:project:1.0}PlainTextSelection", attrib={
                "guid": sel_guid,
                "name": f"Selection-{seg.id}",
                "startPosition": str(seg.start_char),
                "endPosition": str(seg.end_char),
                "creatingUser": master_user_guid
            })

            if hasattr(seg, 'created_at') and seg.created_at:
                sel_elem.set("creationDateTime", seg.created_at.strftime("%Y-%m-%dT%H:%M:%SZ"))

            coding_elem = ET.SubElement(sel_elem, "{urn:QDA-XML:project:1.0}Coding", attrib={
                "guid": generate_guid("coding", seg.id),
                "creatingUser": master_user_guid
            })

            if hasattr(seg, 'created_at') and seg.created_at:
                coding_elem.set("creationDateTime", seg.created_at.strftime("%Y-%m-%dT%H:%M:%SZ"))

            ET.SubElement(coding_elem, "{urn:QDA-XML:project:1.0}CodeRef", attrib={
                "targetGUID": code_guid_map[seg.code_id]
            })

            if seg.id in memos_by_segment:
                for m in memos_by_segment[seg.id]:
                    ET.SubElement(sel_elem, "{urn:QDA-XML:project:1.0}NoteRef", attrib={
                        "targetGUID": generate_guid("memo", m.id)
                    })

    if all_memos:
        notes_elem = ET.SubElement(root, "{urn:QDA-XML:project:1.0}Notes")
        for m in all_memos:
            memo_guid = generate_guid("memo", m.id)
            
            preview_text = (m.text[:47] + '...') if len(m.text) > 50 else m.text
            preview_text = preview_text.replace('\n', ' ')

            note_elem = ET.SubElement(notes_elem, "{urn:QDA-XML:project:1.0}Note", attrib={
                "guid": memo_guid,
                "name": preview_text,
                "creatingUser": master_user_guid
            })
            
            content_elem = ET.SubElement(note_elem, "{urn:QDA-XML:project:1.0}PlainTextContent")
            content_elem.text = m.text

    # Sets
    folders = db.query(models.DocumentFolder).filter(models.DocumentFolder.project_id == project_id).all()
    if folders:
        sets_elem = ET.SubElement(root, "{urn:QDA-XML:project:1.0}Sets")
        for folder in folders:
            set_elem = ET.SubElement(sets_elem, "{urn:QDA-XML:project:1.0}Set", attrib={
                "guid": generate_guid("folder", folder.id),
                "name": folder.name
            })
            for doc in folder.documents:
                ET.SubElement(set_elem, "{urn:QDA-XML:project:1.0}MemberSource", attrib={
                    "targetGUID": generate_guid("doc", doc.id)
                })
            child_folders = [f for f in folders if getattr(f, 'parent_id', None) == folder.id]
            for child in child_folders:
                ET.SubElement(set_elem, "{urn:QDA-XML:project:1.0}MemberSet", attrib={
                    "targetGUID": generate_guid("folder", child.id)
                })

    # Description
    if project.description:
        desc = ET.SubElement(root, "{urn:QDA-XML:project:1.0}Description")
        desc.text = project.description

    for m in project_memos:
        ET.SubElement(root, "{urn:QDA-XML:project:1.0}NoteRef", attrib={
            "targetGUID": generate_guid("memo", m.id)
        })

    # === Compile to ZIP (QDPX) ===
    xml_str = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    safe_filename = project.name.replace(" ", "_")
    
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        zip_file.writestr("project.qde", xml_str)
        
        for filepath, content in zip_files_to_write.items():
            zip_file.writestr(filepath, content.encode('utf-8'))

    zip_bytes = zip_buffer.getvalue()

    return Response(
        content=zip_bytes, 
        media_type="application/zip", 
        headers={
            "Content-Disposition": f'attachment; filename="{safe_filename}.qdpx"'
        }
    )


async def import_refi_xml(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.endswith('.qdpx'):
        raise HTTPException(status_code=400, detail="File must be a .qdpx package")

    try:
        # Unzip the file in memory
        content = await file.read()
        zip_ref = zipfile.ZipFile(io.BytesIO(content))
        
        # Read the XML blueprint
        xml_data = zip_ref.read("project.qde")
        root = ET.fromstring(xml_data)

        #easier parsing
        for elem in root.iter():
            if '}' in elem.tag:
                elem.tag = elem.tag.split('}', 1)[1]

        # Create the Project
        base_name = root.attrib.get("name", "Imported Project")
        proj_name = base_name
        
        #prevent duplicate names
        counter = 1
        while db.query(models.Project).filter(models.Project.name == proj_name).first() is not None:
            proj_name = f"{base_name} ({counter})"
            counter += 1

        desc_elem = root.find("Description")
        proj_desc = desc_elem.text if desc_elem is not None else None

        new_project = models.Project(name=proj_name, description=proj_desc)
        db.add(new_project)
        db.commit()
        db.refresh(new_project)

        #get segment and code memos
        notes_dict = {}
        for note_elem in root.findall(".//Note"): 
            guid = next((v for k, v in note_elem.attrib.items() if k.lower() == "guid"), None)
            text_content = ""

            target_path = next((v for k, v in note_elem.attrib.items() if k.lower() in ["plaintextpath", "richtextpath"]), None)

            if target_path and "internal://" in target_path.lower():
                raw_filename = target_path.split("/")[-1]
                zip_path = next((name for name in zip_ref.namelist() if name.endswith(raw_filename)), None)
                
                if zip_path:
                    raw_bytes = zip_ref.read(zip_path)
                    if zip_path.lower().endswith(".docx"):
                        from docx import Document as DocxDocument
                        doc_obj = DocxDocument(io.BytesIO(raw_bytes))
                        text_content = "\n".join([p.text for p in doc_obj.paragraphs]).strip()
                    else:
                        try:
                            text_content = raw_bytes.decode('utf-8-sig').strip()
                        except UnicodeDecodeError:
                            text_content = raw_bytes.decode('utf-8', errors='ignore').strip()
            if not text_content:
                pt_elem = note_elem.find(".//PlainTextContent")
                if pt_elem is not None and pt_elem.text:
                    text_content = pt_elem.text.strip()
                else: 
                    text_content = "".join(note_elem.itertext()).strip()
                    
            if text_content and guid:
                notes_dict[guid.lower()] = text_content
        
        # Extract Project-Level Memos
        for note_ref in root.findall("./NoteRef"): 
            target_guid = next((v for k, v in note_ref.attrib.items() if k.lower() == "targetguid"), None)
            if target_guid in notes_dict:
                db.add(models.Memo(text=notes_dict[target_guid], target_type="project", target_id=new_project.id))

        # Extract Codes 
        guid_to_code_id = {} # Maps XML GUID to SQLite ID 
        def process_code_elem(code_elem, parent_db_id=None):
            guid = code_elem.attrib.get("guid")
            name = code_elem.attrib.get("name")
            color = code_elem.attrib.get("color", "#646cff")

            date_str = code_elem.attrib.get("modified") or code_elem.attrib.get("created")
            
            c_desc_elem = code_elem.find("Description")
            imported_description = c_desc_elem.text if c_desc_elem is not None else None

            new_code = models.Code(
                project_id=new_project.id,
                name=name,
                color=color,
                parent_id=parent_db_id
            )
            parsed_date = parse_refi_date(date_str)
            if parsed_date:
                new_code.created_at = parsed_date

            db.add(new_code)
            db.flush() # doesn't commit the whole transaction
            
            guid_to_code_id[guid] = new_code.id

            if imported_description:
                db.add(models.Memo(text=imported_description, target_type="code", target_id=new_code.id))

            # Extract Code-Level Memos
            for note_ref in code_elem.findall("./NoteRef"):
                t_guid = next((v for k, v in note_ref.attrib.items() if k.lower() == "targetguid"), None)
                if t_guid and t_guid.lower() in notes_dict:
                    db.add(models.Memo(text=notes_dict[t_guid.lower()], target_type="code", target_id=new_code.id))

            #child nodes
            for child_elem in code_elem.findall("./Code"):
                process_code_elem(child_elem, new_code.id)

        codes_container = root.find(".//Codes")
        if codes_container is not None:
            for top_level_code in codes_container.findall("./Code"):
                process_code_elem(top_level_code, None)

        db.commit()

        # Extract Documents and Segments
        for source_elem in root.findall(".//TextSource"):
            doc_guid = source_elem.attrib.get("guid")
            doc_name = source_elem.attrib.get("name")
            
            # Get the text content
            doc_content = ""
            plain_text_path = source_elem.attrib.get("plainTextPath")
            
            if plain_text_path and plain_text_path.startswith("internal://"):
                raw_filename = plain_text_path.split("/")[-1]
                
                # Search the ZIP file to find exactly where s the file
                zip_path = None
                for name in zip_ref.namelist():
                    if name.endswith(raw_filename):
                        zip_path = name
                        break
                
                if zip_path:
                    try:
                        raw_bytes = zip_ref.read(zip_path)
                        try:
                            doc_content = raw_bytes.decode('utf-8-sig') 
                        except UnicodeDecodeError:
                            doc_content = raw_bytes.decode('utf-8', errors='ignore')
                    except Exception as e:
                        print(f"Failed to read file {zip_path} from ZIP: {e}")
            else:
                # Fallback to embedded text if the software didn't zip a physical file
                pt_elem = source_elem.find("PlainTextContent")
                if pt_elem is not None and pt_elem.text:
                    doc_content = pt_elem.text

            new_doc = models.Document(
                project_id=new_project.id,
                filename=doc_name,
                content=doc_content,
                type="txt"
            )
            db.add(new_doc)
            db.commit()
            db.refresh(new_doc)

            doc_guid = next((v for k, v in source_elem.attrib.items() if k.lower() == "guid"), None)
            if doc_guid:
                if 'guid_to_doc_id' not in locals():
                    guid_to_doc_id = {}
                guid_to_doc_id[doc_guid.lower()] = new_doc.id

            # Extract the coded segments for this document
            for sel_elem in source_elem.findall(".//PlainTextSelection"):
                start_pos = int(sel_elem.attrib.get("startPosition", 0))
                end_pos = int(sel_elem.attrib.get("endPosition", 0))
                
                for coding_elem in sel_elem.findall(".//Coding"):
                    code_ref = coding_elem.find("CodeRef")
                    if code_ref is not None:
                        target_guid = code_ref.attrib.get("targetGUID")
                        
                        # Translate the GUID to database ID
                        actual_code_id = guid_to_code_id.get(target_guid)
                        
                        if actual_code_id:
                            # Slice the string to get the exact highlighted text
                            segment_text = doc_content[start_pos:end_pos]

                            date_str = sel_elem.attrib.get("modified") or sel_elem.attrib.get("created")
                            if not date_str and coding_elem is not None:
                                date_str = coding_elem.attrib.get("modified") or coding_elem.attrib.get("created")
                            
                            new_segment = models.Segment(
                                document_id=new_doc.id,
                                code_id=actual_code_id,
                                start_char=start_pos,
                                end_char=end_pos,
                                content=segment_text
                            )
                            parsed_date = parse_refi_date(date_str)
                            if parsed_date:
                                new_segment.created_at = parsed_date

                            db.add(new_segment)
                            db.flush()

                            for note_ref in sel_elem.findall(".//NoteRef"):
                                t_guid = next((v for k, v in note_ref.attrib.items() if k.lower() == "targetguid"), None)
                                if t_guid and t_guid.lower() in notes_dict:
                                    db.add(models.Memo(text=notes_dict[t_guid.lower()], target_type="segment", target_id=new_segment.id))


        if 'guid_to_doc_id' in locals():
            guid_to_folder_id = {}
            
            # create all folders first so we have their IDs
            for set_elem in root.findall(".//Set"):
                set_guid = set_elem.attrib.get("guid") or str(uuid.uuid4())
                set_name = set_elem.attrib.get("name", "Imported Set")
                
                new_folder = models.DocumentFolder(name=set_name, project_id=new_project.id)
                db.add(new_folder)
                db.flush()
                guid_to_folder_id[set_guid.lower()] = new_folder.id
                
                # Assign documents to this folder
                for ms in set_elem.findall("./MemberSource"):
                    t_guid = next((v for k, v in ms.attrib.items() if k.lower() == "targetguid"), None)
                    if t_guid and t_guid.lower() in guid_to_doc_id:
                        doc_id = guid_to_doc_id[t_guid.lower()]
                        db.query(models.Document).filter(models.Document.id == doc_id).update({"folder_id": new_folder.id})
            
            # establish nested sub-folder relationships (MemberSet)
            for set_elem in root.findall(".//Set"):
                parent_guid = set_elem.attrib.get("guid")
                if not parent_guid or parent_guid.lower() not in guid_to_folder_id:
                    continue
                parent_id = guid_to_folder_id[parent_guid.lower()]
                
                for ms in set_elem.findall("./MemberSet"):
                    child_guid = next((v for k, v in ms.attrib.items() if k.lower() == "targetguid"), None)
                    if child_guid and child_guid.lower() in guid_to_folder_id:
                        child_id = guid_to_folder_id[child_guid.lower()]
                        # Update the child folder to point to this parent
                        db.query(models.DocumentFolder).filter(models.DocumentFolder.id == child_id).update({"parent_id": parent_id})
                        
        db.commit()
        return new_project
    
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="Invalid QDPX package (Not a valid ZIP file)")
    except Exception as e:
        db.rollback() # If anything fails, cancel the database transaction
        print(f"Import Error: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to process QDPX: {str(e)}")