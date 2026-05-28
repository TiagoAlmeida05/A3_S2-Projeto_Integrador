from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from database import get_db
import schemas
from repositories.code_repo import CodeRepository

#codes export imports
import io
import urllib.parse
from fastapi.responses import StreamingResponse
from docx import Document
from docx.shared import Pt, RGBColor
import models

router = APIRouter(prefix="/projects/{project_id}/codes", tags=["Codes"])

def get_code_repo(db: Session = Depends(get_db)):
    return CodeRepository(db)

@router.get("/")
def get_project_codes(project_id: int, repo: CodeRepository = Depends(get_code_repo)):
    codes = repo.get_by_project(project_id)
    return [{
        "id": c.id, "name": c.name, "color": c.color, "description": c.description,
        "project_id": c.project_id, "parent_id": c.parent_id, "order_index": c.order_index,
        "frequency": len(c.segments) 
    } for c in codes]

@router.post("/")
def create_code(project_id: int, code: schemas.CodeCreate, repo: CodeRepository = Depends(get_code_repo)):
    return repo.create(project_id, code)

@router.put("/reorder")
def reorder_codes(project_id: int, reorder_request: schemas.CodeReorderRequest, repo: CodeRepository = Depends(get_code_repo)):
    repo.reorder(project_id, reorder_request.codes)
    return {"message": "Codes reordered successfully"}

@router.put("/{code_id}")
def update_code(project_id: int, code_id: int, code_update: schemas.CodeUpdate, repo: CodeRepository = Depends(get_code_repo)):
    code = repo.update(project_id, code_id, code_update)
    if not code:
        raise HTTPException(status_code=404, detail="Code not found")
    return code

@router.delete("/{code_id}")
def delete_code(project_id: int, code_id: int, repo: CodeRepository = Depends(get_code_repo)):
    success = repo.delete(project_id, code_id)
    if not success:
        raise HTTPException(status_code=404, detail="Code not found")
    return {"message": "Code deleted successfully"}

@router.get("/export/docx")
def export_codebook_docx(project_id: int, db: Session = Depends(get_db)):
    # fetch project
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # fetch all codes and memos for this project
    codes = db.query(models.Code).filter(models.Code.project_id == project_id).order_by(models.Code.order_index).all()
    
    code_ids = [c.id for c in codes]
    memos = db.query(models.Memo).filter(
        models.Memo.target_type == "code", 
        models.Memo.target_id.in_(code_ids)
    ).all()
    
    # group memos by code_id
    memo_dict = {}
    for m in memos:
        if m.target_id not in memo_dict:
            memo_dict[m.target_id] = []
        memo_dict[m.target_id].append(m.text)

    # build the hierarchical tree recursively
    def build_tree(parent_id=None, depth=0):
        tree = []
        for c in [c for c in codes if c.parent_id == parent_id]:
            tree.append((c, depth))
            tree.extend(build_tree(c.id, depth + 1))
        return tree

    ordered_codes = build_tree()

    # initialize the Word Document
    doc = Document()
    doc.add_heading(f"Codebook: {project.name}", 0)
    doc.add_paragraph(f"Exported from jUPiter QDA on {project.last_accessed.strftime('%B %d, %Y')}")

    # populate the document
    for code, depth in ordered_codes:
        # heading levels 1 to 4
        level = min(depth + 1, 4) 
        heading = doc.add_heading(level=level)
        
        # create an indent string (4 spaces per depth level)
        indent_prefix = "    " * depth 
        
        # add the spaces before the code name
        run = heading.add_run(f"{indent_prefix}{code.name}")
            
        if code.id in memo_dict:
            for memo_text in memo_dict[code.id]:
                p = doc.add_paragraph(style='List Bullet')
                # use Word's native left indent for bullets so the actual bullet dot moves over!
                p.paragraph_format.left_indent = Pt(24 * (depth + 1))
                p.add_run("Memo: ").bold = True
                p.add_run(memo_text)

    # save to a virtual file in memory
    mem_stream = io.BytesIO()
    doc.save(mem_stream)
    mem_stream.seek(0)
    
    # safely encode the filename
    safe_filename = urllib.parse.quote(f"Codebook_{project.name}.docx")
    
    return StreamingResponse(
        mem_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=utf-8''{safe_filename}"}
    )